"""Streamlit UI: upload video -> GPU transcript + speaker detection -> label speakers -> readable transcript.

Run: streamlit run app.py
"""
import gc
import io
import os
import tempfile
import time
from pathlib import Path

import streamlit as st

import diarize
from transcribe import add_cuda_dll_paths, fmt_ts, load_model

MODELS = ["large-v3-turbo", "large-v3", "medium", "small", "base", "distil-large-v3"]
UPLOAD_DIR = Path(__file__).resolve().parent / ".tmp"  # keep temp files on D:
PARA_GAP_S = 2.0      # silence that starts a new paragraph
PARA_MAX_CHARS = 600  # force a paragraph break after this many characters

st.set_page_config(page_title="Video Transcript", page_icon="📝", layout="wide")


def _sweep_stale_temp(max_age_h: float = 6.0):
    """A run killed mid-way (Ctrl+C, crash) leaves its multi-GB temp upload behind; remove old ones."""
    if not UPLOAD_DIR.is_dir():
        return
    for f in UPLOAD_DIR.glob("tmp*"):
        try:
            if f.is_file() and time.time() - f.stat().st_mtime > max_age_h * 3600:
                f.unlink()
        except OSError:
            pass


_sweep_stale_temp()


@st.cache_resource(show_spinner="Loading model on GPU…")
def get_model(name: str, device: str):
    add_cuda_dll_paths()
    return load_model(name, device)


def hms(t: float) -> str:
    return fmt_ts(t).split(",")[0]


def to_paragraphs(segs: list[dict]) -> list[dict]:
    paras, cur = [], None
    for s in segs:
        same = cur and s.get("spk") == cur.get("spk")
        if same and s["start"] - cur["end"] < PARA_GAP_S and len(cur["text"]) < PARA_MAX_CHARS:
            cur["text"] += " " + s["text"]
            cur["end"] = s["end"]
        else:
            cur = {"start": s["start"], "end": s["end"], "text": s["text"], "spk": s.get("spk")}
            paras.append(cur)
    return paras


def who(p: dict, names: dict) -> str:
    return names.get(p["spk"], "") if p.get("spk") is not None else ""


def readable_text(paras, names, timestamps) -> str:
    out = []
    for p in paras:
        head = (f"[{hms(p['start'])}] " if timestamps else "") + (f"{who(p, names)}: " if who(p, names) else "")
        out.append(head + p["text"])
    return "\n\n".join(out)


def build_srt(segs, names) -> str:
    out = []
    for i, s in enumerate(segs, 1):
        n = names.get(s.get("spk"), "")
        out.append(f"{i}\n{fmt_ts(s['start'])} --> {fmt_ts(s['end'])}\n{n + ': ' if n else ''}{s['text']}\n\n")
    return "".join(out)


def build_docx(title, paras, names, timestamps) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=1)
    for p in paras:
        para = doc.add_paragraph()
        if timestamps:
            para.add_run(f"[{hms(p['start'])}] ")
        if who(p, names):
            para.add_run(who(p, names) + ": ").bold = True
        para.add_run(p["text"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- sidebar
with st.sidebar:
    st.header("Settings")
    model_name = st.selectbox("Model", MODELS, help="large-v3-turbo: best speed/accuracy on 8 GB GPU")
    lang = st.text_input("Language code (blank = auto)", "").strip() or None
    task = st.radio("Task", ["transcribe", "translate"], format_func=lambda x: {
        "transcribe": "Transcribe (original language)", "translate": "Translate to English"}[x])
    device = st.radio("Device", ["cuda", "cpu"])
    vad = st.checkbox("Skip silence (VAD)", True)
    timestamps = st.checkbox("Show timestamps", True)
    st.divider()
    detect = st.checkbox("Detect speakers", True)
    engine = st.radio("Speaker engine", ["sortformer", "resnet"], disabled=not detect, format_func=lambda x: {
        "sortformer": "NVIDIA Sortformer (GPU, up to 4 speakers)",
        "resnet": "ResNet293 voice clustering (CPU, any number)"}[x])
    n_spk = st.number_input("Number of speakers (0 = auto)", 0, 20, 0, disabled=not detect or engine != "resnet",
                            help="ResNet only. Set it if you know it; auto-detect can over/under-count.")

st.title("📝 Video Transcript")
up = st.file_uploader("Upload a video or audio file",
                      type=["mp4", "mkv", "mov", "avi", "webm", "m4v", "mp3", "wav", "m4a", "flac", "ogg"])
local = st.text_input("…or path of a file already on this PC (no upload; best for big recordings)", "",
                      placeholder=r"D:\Recordings\meeting.mp4").strip().strip('"')
src_name = Path(up.name).stem if up else (Path(local).stem if local else "")

# ---------------------------------------------------------------- step 1: process
if (up or local) and st.button("Transcribe", type="primary"):
    st.session_state.pop("result", None)
    st.session_state.pop("names", None)
    UPLOAD_DIR.mkdir(exist_ok=True)
    status = st.status("Starting…", expanded=True)
    tmp_path = wav_path = samples = None
    try:
        if up:
            status.write("1/4 Saving upload to D: …")
            with tempfile.NamedTemporaryFile(delete=False, suffix=Path(up.name).suffix, dir=UPLOAD_DIR) as tmp:
                tmp.write(up.getbuffer())
                tmp_path = tmp.name
            src = tmp_path
        else:
            if not Path(local).is_file():
                raise FileNotFoundError(f"File not found: {local}")
            src = local
        status.write("2/4 Extracting audio …")
        with tempfile.NamedTemporaryFile(delete=False, suffix=".wav", dir=UPLOAD_DIR) as tmp:
            wav_path = tmp.name
        diarize.extract_audio(src, wav_path)

        status.write("3/4 Loading model & transcribing …")
        model, mode = get_model(model_name, device)
        t0 = time.time()
        gen, info = model.transcribe(
            wav_path, language=lang, task=task, beam_size=5, vad_filter=vad,
            vad_parameters={"min_silence_duration_ms": 500}, word_timestamps=detect,
        )
        bar = st.progress(0.0, text=f"Transcribing on {mode}…")
        segs, words = [], []
        for s in gen:
            segs.append({"start": s.start, "end": s.end, "text": s.text.strip()})
            words += [{"start": w.start, "end": w.end, "word": w.word} for w in (s.words or [])]
            if info.duration:
                bar.progress(min(s.end / info.duration, 1.0), text=f"{hms(s.end)} / {hms(info.duration)}")
        bar.empty()

        speakers = {}
        if detect and words:
            status.write("4/4 Detecting speakers (first run downloads the speaker model) …")
            samples = diarize.load_wav(wav_path)  # memory-mapped from D:, not loaded into RAM
            bar = st.progress(0.0, text="Detecting speakers…")
            prog = lambda f: bar.progress(min(f, 1.0))  # noqa: E731
            if engine == "sortformer":
                import sortformer
                dsegs, secs, turns = sortformer.run(wav_path, words, device, progress=prog)
            else:
                dsegs, secs, turns = diarize.run(samples, words, int(n_spk), progress=prog)
            bar.empty()
            if dsegs:
                segs = dsegs
                clips = diarize.speaker_samples(samples, turns)
                speakers = {k: {"secs": secs[k], "clip": clips[k]} for k in sorted(secs) if k in clips}

        status.update(label="Done", state="complete", expanded=False)
        st.session_state["result"] = {
            "name": src_name, "segs": segs, "lang": info.language, "speakers": speakers,
            "duration": info.duration, "elapsed": time.time() - t0, "mode": mode,
        }
        st.session_state["labelled"] = not speakers  # no speakers -> skip labelling step
    except Exception as e:
        status.update(label="Failed", state="error")
        st.error(f"Failed: {e}")
    finally:
        samples = None
        gc.collect()  # release the memory-mapped wav so Windows lets us delete it
        for p in (tmp_path, wav_path):
            try:
                if p:
                    os.remove(p)
            except OSError:
                pass

r = st.session_state.get("result")

# ---------------------------------------------------------------- step 2: label speakers
if r and not st.session_state.get("labelled"):
    st.subheader(f"{len(r['speakers'])} speaker(s) detected — listen and name them")
    with st.form("labels"):
        cols = st.columns(min(len(r["speakers"]), 4))
        for i, (k, sp) in enumerate(r["speakers"].items()):
            with cols[i % len(cols)]:
                st.markdown(f"**Speaker {k + 1}** · {hms(sp['secs'])} talk time")
                st.audio(sp["clip"], format="audio/wav")
                st.text_input("Name", value=f"Speaker {k + 1}", key=f"name_{k}")
        if st.form_submit_button("Generate transcript", type="primary"):
            st.session_state["names"] = {k: st.session_state[f"name_{k}"].strip() or f"Speaker {k + 1}"
                                         for k in r["speakers"]}
            st.session_state["labelled"] = True
            st.rerun()

# ---------------------------------------------------------------- step 3: final transcript
elif r:
    names = st.session_state.get("names", {})
    paras = to_paragraphs(r["segs"])
    st.success(
        f"Done in {r['elapsed']:.0f}s on {r['mode']} · language: {r['lang']} · "
        f"{r['duration'] / max(r['elapsed'], 1e-6):.1f}× realtime"
        + (f" · {len(names)} speakers" if names else "")
    )
    if names and st.button("✏ Edit speaker names"):
        st.session_state["labelled"] = False
        st.rerun()

    txt = readable_text(paras, names, timestamps)
    base = r["name"]
    c1, c2, c3 = st.columns(3)
    c1.download_button("⬇ Transcript (.txt)", txt, f"{base}.txt", "text/plain", use_container_width=True)
    c2.download_button("⬇ Subtitles (.srt)", build_srt(r["segs"], names), f"{base}.srt", "text/plain",
                       use_container_width=True)
    try:
        c3.download_button("⬇ Word (.docx)", build_docx(base, paras, names, timestamps), f"{base}.docx",
                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                           use_container_width=True)
    except ImportError:
        c3.caption("pip install python-docx for .docx")

    st.divider()
    for p in paras:
        ts = f"`{hms(p['start'])}` " if timestamps else ""
        nm = f"**{who(p, names)}** " if who(p, names) else ""
        st.markdown(f"{ts}{nm}{p['text']}")
